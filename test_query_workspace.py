"""Run against the isolated demo DB. No live LLM calls or cloud writes."""
from datetime import date
from pathlib import Path
import tempfile
import uuid
from unittest.mock import patch
from openpyxl import load_workbook
import pipeline as p
import query_workspace as w


def check():
    with p.connect() as conn,tempfile.TemporaryDirectory() as tmp:
        day=date(2026,9,29)
        schema=w.tables(conn)
        assert 'mvp_raw_daily' in schema and schema['mvp_raw_daily'][0]['name']=='report_date'
        assert schema['mvp_raw_daily'][0]['primary']
        model=w.schema(conn)
        assert any(r['kind']=='derived' for r in model['relationships'])
        # Verify real FK metadata and views without altering existing analytics tables.
        base='test_er_'+uuid.uuid4().hex[:8];parent,child,view=base+'_p',base+'_c',base+'_v'
        try:
            with conn.cursor() as c:
                c.execute(f'CREATE TABLE `{parent}` (id INT PRIMARY KEY)')
                c.execute(f'CREATE TABLE `{child}` (id INT PRIMARY KEY, parent_id INT, FOREIGN KEY (parent_id) REFERENCES `{parent}`(id))')
                c.execute(f'CREATE VIEW `{view}` AS SELECT id FROM `{parent}`')
            model=w.schema(conn)
            assert any(r['kind']=='foreign_key' and r['from']==child and r['columns']==[{'from':'parent_id','to':'id'}] for r in model['relationships'])
            assert any(t['name']==view and t['kind']=='VIEW' for t in model['tables'])
            assert not w.preview(conn,view,day)['rows']
        finally:
            with conn.cursor() as c:
                c.execute(f'DROP VIEW IF EXISTS `{view}`');c.execute(f'DROP TABLE IF EXISTS `{child}`');c.execute(f'DROP TABLE IF EXISTS `{parent}`')
        rows=p.mock_rows((day,day.replace(day=22)))
        p.acquire_lock(conn)
        try:p.load_transform(conn,rows,(day.replace(day=22),day));conn.commit()
        finally:
            with conn.cursor() as c:c.execute('SELECT RELEASE_LOCK(%s)',(p.lock_name(conn),))
        for sql in ('DELETE FROM mvp_raw_daily','SELECT 1; DELETE FROM mvp_raw_daily',"SELECT LOAD_FILE('/etc/passwd')",'SELECT SLEEP(10)','SELECT 1 INTO OUTFILE \'/tmp/output\'','SELECT /* hint */ 1'):
            try:w.validate_sql(sql)
            except ValueError:pass
            else:raise AssertionError('Unsafe SQL accepted')
        result=w.execute(conn,'SELECT report_date, SUM(revenue_krw) AS revenue FROM mvp_raw_daily WHERE report_date=%(report_date)s GROUP BY report_date',w.parameters(day))
        assert int(result['rows'][0]['revenue'])==sum(r['revenue_krw'] for r in rows if r['date']==day.isoformat())
        assert len(w.execute(conn,'SELECT * FROM mvp_raw_daily',w.parameters(day),2)['rows'])==2
        assert w.execute(conn,'SELECT * FROM mvp_raw_daily',w.parameters(day),2)['truncated']
        preview=w.preview(conn,'mvp_raw_daily',day)
        assert len(preview['rows'])==20
        assert preview['rows'][0]['report_date']==max(r['report_date'] for r in w.execute(conn,'SELECT report_date FROM mvp_raw_daily',w.parameters(day),5000)['rows'])
        ordered=w.execute(conn,'WITH chosen AS (SELECT report_date FROM mvp_raw_daily) SELECT report_date FROM chosen ORDER BY report_date DESC LIMIT 3',w.parameters(day))
        assert len(ordered['rows'])==3 and ordered['rows'][0]['report_date']==preview['rows'][0]['report_date']
        for name in ('missing_table','mvp_raw_daily` WHERE 1=1','information_schema.COLUMNS'):
            try:w.preview(conn,name,day)
            except ValueError:pass
            else:raise AssertionError('Unknown table accepted')
        builtin=w.selected('builtin:history-channel')
        assert len(w.execute(conn,builtin['sql'],w.parameters(day))['rows'])==5
        table=w.selected('table:mvp_history_daily',conn)
        assert len(w.execute(conn,table['sql'],w.parameters(day))['rows'])==5
        root=Path(tmp)
        with patch.object(w,'AGGREGATION',root/'aggregation.json'):
            default=w.aggregation_settings()['default_sql']
            custom=default.replace('SUM(revenue_krw)','SUM(revenue_krw)+0')
            assert w.preview_aggregation(conn,custom,day)['validated']
            for bad in (default.replace(' AS visits',' AS other'),default.replace('SUM(revenue_krw)','SUM(revenue_krw)+1'),default.replace('IN (%(comparison_date)s, %(report_date)s)','= %(report_date)s')):
                try:w.preview_aggregation(conn,bad,day)
                except ValueError:pass
                else:raise AssertionError('Invalid aggregation accepted')
            w.save_aggregation(custom)
            assert w.aggregation_settings()['custom']
            conn.begin()
            with patch.object(w,'execute',wraps=w.execute) as execute:
                p.load_transform(conn,rows,(day.replace(day=22),day))
                assert any(call.args[1]==custom and call.kwargs.get('own_transaction') is False for call in execute.call_args_list)
            conn.commit()
            # Failed custom aggregation must roll back new raw data and preserve history.
            snapshots={name:w.execute(conn,'SELECT * FROM '+name+' ORDER BY report_date,channel'+(',device' if name=='mvp_raw_daily' else ''),w.parameters(day),5000)['rows'] for name in ('mvp_raw_daily','mvp_history_daily')}
            conn.begin()
            try:p.load_transform(conn,[dict(r,revenue_krw=r['revenue_krw']+1) for r in rows],(day.replace(day=22),day),default.replace('SUM(revenue_krw)','SUM(revenue_krw)+1'))
            except ValueError:conn.rollback()
            else:raise AssertionError('Failed aggregation wrote to DB')
            for name,before in snapshots.items():
                assert w.execute(conn,'SELECT * FROM '+name+' ORDER BY report_date,channel'+(',device' if name=='mvp_raw_daily' else ''),w.parameters(day),5000)['rows']==before
        with patch.object(w,'PIPELINE_STORE',root/'pipeline_sql.json'):
            files=w.pipeline_contents()
            snapshots={name:w.execute(conn,'SELECT * FROM '+name+' ORDER BY report_date,channel'+(',device' if name=='mvp_raw_daily' else ''),w.parameters(day),5000)['rows'] for name in ('mvp_raw_daily','mvp_history_daily')}
            for name,sql in files.items():
                assert w.preview_pipeline(conn,name,sql,day)['validated']
            # A broad DELETE with valid bindings must still fail and roll back outside-date changes.
            bad=files['load_daily.sql'].replace('report_date IN (%s, %s)','report_date IN (%s, %s) OR 1=1')
            try:w.preview_pipeline(conn,'load_daily.sql',bad,day)
            except ValueError as error:assert '대상 날짜 밖' in str(error)
            else:raise AssertionError('Out-of-scope write accepted')
            bad=files['report_daily.sql'].replace('channel, device, visits','channel, device, visits + 1 AS visits',1)
            try:w.preview_pipeline(conn,'report_daily.sql',bad,day)
            except ValueError:pass
            else:raise AssertionError('Broken report output accepted')
            for name,before in snapshots.items():
                assert w.execute(conn,'SELECT * FROM '+name+' ORDER BY report_date,channel'+(',device' if name=='mvp_raw_daily' else ''),w.parameters(day),5000)['rows']==before
            for bad in (files['load_daily.sql'].replace('mvp_raw_daily','other_table'),files['load_daily.sql'].replace('mvp_raw_daily WHERE','mvp_raw_daily , other_table WHERE'),'DROP TABLE mvp_raw_daily'):
                try:w.validate_pipeline_sql('load_daily.sql',bad)
                except ValueError:pass
                else:raise AssertionError('Unsupported write accepted')
            custom=files['report_daily.sql'].replace('ORDER BY report_date, channel, device','ORDER BY report_date DESC, channel, device')
            w.save_pipeline_sql('report_daily.sql',custom)
            conn.begin()
            actual=p.report_data(conn,(day.replace(day=22),day),rows)[0]
            assert actual[0]['date']==day.isoformat()
            conn.rollback()
        with patch.object(w,'STORE',root/'queries.json'),patch.dict(p.os.environ,{'REPORT_DIR':str(root),'LLM_PROVIDER':'none'}):
            q=w.save_query('매출 조회',"SELECT '=1+1' AS note, SUM(revenue_krw) AS revenue FROM mvp_raw_daily WHERE report_date=%(report_date)s")
            assert w.selected(q['id'])['title']=='매출 조회'
            assert all(not x.get('builtin') for x in __import__('json').loads(w.STORE.read_text()))
            w.save_query('수정',builtin['sql'],builtin['id'])
            assert w.selected(builtin['id'])['title']=='수정' and w.selected(builtin['id'])['builtin']
            with patch.object(w.providers,'summarize',return_value={'provider':'none','model':'','status':'disabled','text':'=1+1'}):
                output=w.report(conn,day,q['id']);book=load_workbook(output)
                assert book['조회 데이터']['A2'].data_type=='s' and book['조회 데이터']['B2'].data_type=='n'
                assert book['AI 분석']['A3'].data_type=='s';book.close()
                assert w.report(conn,day,q['id'])==output
                assert w.report(conn,day,q['id'],True)!=output and output.exists()
                tomorrow=day.replace(day=30)
                with patch.object(p,'extract',return_value=p.mock_rows((tomorrow.replace(day=23),tomorrow))) as extract:
                    assert w.report(conn,tomorrow,q['id'],refresh=True).exists()
                    extract.assert_called_once_with((tomorrow.replace(day=23),tomorrow),None)
                empty=w.save_query('빈 조회','SELECT * FROM mvp_raw_daily WHERE 1=0')
                try:w.report(conn,day,empty['id'])
                except ValueError:pass
                else:raise AssertionError('Empty report published')
    print('PASS: schema discovery, SELECT limits, SQL rejection, saved query, Excel numbers/formula safety and version preservation')


if __name__=='__main__':check()
