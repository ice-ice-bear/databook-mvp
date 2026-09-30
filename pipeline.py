"""Local MariaDB daily report MVP. Selectable example data and optional LLM analysis."""
import argparse
import csv
import hashlib
import json
import logging
import os
import random
import re
import tempfile
import uuid
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import providers

import pymysql
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent
KST = timezone(timedelta(hours=9), 'KST')
CHANNELS = ('paid_search', 'organic_search', 'direct', 'email', 'social')
DEVICES = ('mobile', 'desktop', 'tablet')
METRICS = ('visits', 'page_views', 'orders', 'revenue_krw')
FIELDS = ('date', 'channel', 'device', *METRICS)
VERSION = '3'
LOG = logging.getLogger('databook')


def sql_statements(name,contents=None):
    import query_workspace as w
    if name in w.PIPELINE_FILES:
        return w.validate_pipeline_sql(name,(contents or w.pipeline_contents())[name])
    return [part.strip() for part in (ROOT/'sql'/name).read_text(encoding='utf-8').split(';') if part.strip()]


def read_bound(conn,sql,params=()):
    import query_workspace as w
    bindings={'bound_'+chr(97+i):value for i,value in enumerate(params)}
    names=iter(bindings)
    sql=re.sub(r'%s',lambda _: '%('+next(names)+')s',sql)
    return w.execute(conn,sql,bindings,own_transaction=False)


def connect():
    db = os.environ.get('DB_NAME', 'databook_mvp_test')
    if not re.fullmatch(r'databook_mvp_[a-z0-9_]+', db):
        raise ValueError('This MVP writes only to a designated databook_mvp_* test database.')
    return pymysql.connect(
        host=os.environ.get('DB_HOST', '127.0.0.1'),
        port=int(os.environ.get('DB_PORT', '3306')), database=db,
        user=os.environ.get('DB_USER', 'databook_mvp'),
        password=os.environ.get('DB_PASSWORD', ''), charset='utf8mb4',
        cursorclass=pymysql.cursors.DictCursor, autocommit=False,
        connect_timeout=10, read_timeout=30, write_timeout=30,
    )


def mock_rows(days):
    """Same day always produces the same data, including overlapping reruns."""
    rows = []
    for day in days:
        rng = random.Random(day.isoformat())
        for ci, channel in enumerate(CHANNELS):
            for device, share, rate in zip(DEVICES, (.65, .30, .05), (.023, .036, .018)):
                volume = (4500, 3600, 3000, 1200, 1900)[ci] * share
                volume *= (.84 if day.weekday() >= 5 else 1) * rng.uniform(.97, 1.03)
                cvr = rate * (1.05, 1, 1.12, 1.30, .65)[ci]
                # A fixed synthetic event; never inferred as a real-world cause.
                if day >= date(2026, 9, 21) and channel == 'paid_search' and device == 'mobile':
                    volume *= .62
                    cvr *= .67
                if day >= date(2026, 9, 21) and channel == 'email':
                    volume *= 1.10
                visits = round(volume)
                orders = round(visits * cvr)
                aov = round(72000 * (1.08 if device == 'desktop' else 1) * rng.uniform(.98, 1.02))
                rows.append(dict(date=day.isoformat(), channel=channel, device=device,
                                 visits=visits, page_views=round(visits*rng.uniform(2.6, 3.8)),
                                 orders=orders, revenue_krw=orders*aov))
    return rows


def extract(days, csv_path=None):
    if csv_path:
        with Path(csv_path).open(encoding='utf-8-sig', newline='') as handle:
            reader = csv.DictReader(handle)
            if reader.fieldnames != list(FIELDS):
                raise ValueError(f'CSV columns must be {FIELDS}')
            rows = [r for r in reader if r['date'] in {d.isoformat() for d in days}]
    elif os.environ.get('DATA_PROVIDER', 'mock') == 'example_api':
        rows = providers.example_rows(days)
    elif os.environ.get('DATA_PROVIDER', 'mock') == 'mock':
        rows = mock_rows(days)
    elif os.environ.get('DATA_PROVIDER') == 'adobe':
        import adobe
        rows = adobe.daily_rows(days)
    else:
        raise ValueError('DATA_PROVIDER must be mock, example_api or adobe; use --csv for CSV')
    return validate(rows, days)


def validate(rows, days):
    expected = {(d.isoformat(), c, v) for d in days for c in CHANNELS for v in DEVICES}
    seen, normalized = set(), []
    for row in rows:
        key = (row['date'], row['channel'], row['device'])
        if key not in expected or key in seen:
            raise ValueError(f'Unexpected or duplicate source key: {key}')
        seen.add(key)
        clean = dict(zip(FIELDS[:3], key))
        for metric in METRICS:
            value = row[metric]
            if isinstance(value, bool) or not re.fullmatch(r'[0-9]{1,15}', str(value)):
                raise ValueError(f'Invalid nonnegative integer: {metric} at {key}')
            clean[metric] = int(value)
        if clean['page_views'] < clean['visits']:
            raise ValueError(f'page_views below visits at {key}')
        normalized.append(clean)
    if seen != expected:
        raise ValueError(f'Incomplete source: expected {len(expected)} rows, received {len(seen)}')
    return sorted(normalized, key=lambda r: (r['date'], r['channel'], r['device']))


def lock_name(conn):
    database = conn.db.decode() if isinstance(conn.db, bytes) else conn.db
    return f'databook-mvp:{database}'


def acquire_lock(conn):
    # ponytail: one DB-wide pipeline lock; split by report dates if throughput requires it.
    with conn.cursor() as cur:
        cur.execute('SELECT GET_LOCK(%s, 0) AS acquired', (lock_name(conn),))
        if cur.fetchone()['acquired'] != 1:
            raise RuntimeError('Another pipeline run is active. Retry after it finishes.')


def init_db(conn):
    acquire_lock(conn)
    try:
        with conn.cursor() as cur:
            for statement in sql_statements('schema.sql'):
                if statement.strip():
                    cur.execute(statement)
        conn.commit()
    finally:
        with conn.cursor() as cur:
            cur.execute('SELECT RELEASE_LOCK(%s)', (lock_name(conn),))


def load_transform(conn, rows, days, aggregation_sql=None,contents=None):
    """Caller owns the transaction: failed transform rolls back raw replacement too."""
    import query_workspace as w
    contents=contents or w.pipeline_contents()
    with conn.cursor() as cur:
        # This transaction guard is independent of the user's editable validation SELECT.
        cur.execute("SELECT TABLE_NAME, ENGINE FROM information_schema.TABLES WHERE TABLE_SCHEMA=DATABASE() AND TABLE_NAME IN ('mvp_raw_daily','mvp_history_daily')")
        engines = cur.fetchall()
        if len(engines) != 2 or any(x['ENGINE'] != 'InnoDB' for x in engines):
            raise RuntimeError('Run init-db first. Both MVP tables must use InnoDB.')
        checks = sql_statements('validate_daily.sql',contents)
        engine_check=read_bound(conn,checks[0])
        if {r.get('TABLE_NAME') for r in engine_check['rows']}!={'mvp_raw_daily','mvp_history_daily'} or any(r.get('ENGINE')!='InnoDB' for r in engine_check['rows']):raise ValueError('파이프라인 스키마 검증은 두 MVP 테이블의 TABLE_NAME·ENGINE을 반환해야 합니다.')
        # ponytail: snapshot both small MVP tables; use staging tables for large datasets.
        outside={}
        for table in ('mvp_raw_daily','mvp_history_daily'):
            cur.execute('SELECT * FROM '+table+' WHERE report_date NOT IN (%s,%s) ORDER BY report_date,channel'+(',device' if table=='mvp_raw_daily' else ''),days)
            outside[table]=cur.fetchall()
        load = sql_statements('load_daily.sql',contents)
        cur.execute('SET STATEMENT max_statement_time=5 FOR '+load[0], days)
        cur.executemany('SET STATEMENT max_statement_time=5 FOR '+load[1],
                        [tuple(r[k] for k in FIELDS) for r in rows])
        LOG.info('load: %s validated rows', len(rows))
        cur.execute('SELECT report_date AS date, channel, device, visits, page_views, orders, revenue_krw FROM mvp_raw_daily WHERE report_date IN (%s,%s)',days)
        validate([dict(r,date=r['date'].isoformat()) for r in cur.fetchall()],days)
        sql=aggregation_sql or sql_statements('aggregate_daily.sql',contents)[0]
        result=w.execute(conn,sql,dict(w.parameters(days[1]),comparison_date=days[0].isoformat()),10,own_transaction=False)
        baseline=analyze([dict(r,report_date=date.fromisoformat(r['date'])) for r in rows],days[1],days[0])
        aggregated=w.validate_aggregation(result,days[1],baseline,days[0])
        transform=sql_statements('transform.sql',contents)
        cur.execute('SET STATEMENT max_statement_time=5 FOR '+transform[0],days)
        cur.executemany('SET STATEMENT max_statement_time=5 FOR '+transform[1],[tuple(r[k] for k in ('report_date','channel',*METRICS)) for r in aggregated])
        LOG.info('transform: channel aggregates written to history')
        checked=read_bound(conn,checks[1],days)
        history = checked['rows']
        w.validate_aggregation(checked,days[1],baseline,days[0])
        for table,before in outside.items():
            cur.execute('SELECT * FROM '+table+' WHERE report_date NOT IN (%s,%s) ORDER BY report_date,channel'+(',device' if table=='mvp_raw_daily' else ''),days)
            if cur.fetchall()!=before:raise ValueError('파이프라인 SQL이 대상 날짜 밖의 데이터를 변경했습니다. 변경을 롤백합니다.')
        for day in days:
            selected = [h for h in history if h['report_date'] == day]
            if len(selected) != len(CHANNELS):
                raise ValueError('Incomplete transformed data')
            for metric in METRICS:
                if sum(h[metric] for h in selected) != sum(r[metric] for r in rows if r['date'] == day.isoformat()):
                    raise ValueError(f'Transform reconciliation failed: {metric}')
        return history


def report_data(conn,days,source,contents=None):
    import query_workspace as w
    results=[]
    for sql in sql_statements('report_daily.sql',contents):
        result=read_bound(conn,sql,days)
        if result['truncated']:raise ValueError('파이프라인 보고서 SQL은 두 날짜의 집계·상세 데이터만 반환하세요.')
        results.append(result)
    try:
        detail=[dict(r,date=r['date'].isoformat() if isinstance(r['date'],date) else r['date']) for r in results[0]['rows']]
        validate(detail,days)
        baseline=analyze([dict(r,report_date=date.fromisoformat(r['date'])) for r in source],days[1],days[0])
        history=w.validate_aggregation(results[1],days[1],baseline,days[0])
        devices=results[2]['rows']
        expected={(d,v) for d in days for v in DEVICES}
        if len(devices)!=len(expected) or {(r['report_date'],r['device']) for r in devices}!=expected:raise ValueError()
        if any(int(r[k])!=r[k] or r[k]<0 for r in devices for k in METRICS):raise ValueError()
        for d in days:
            for metric in METRICS:
                total=baseline[d.isoformat()][metric]
                if sum(r[metric] for r in devices if r['report_date']==d)!=total or sum(r[metric] for r in detail if r['date']==d.isoformat())!=total:raise ValueError()
    except (KeyError,TypeError,ValueError):raise ValueError('파이프라인 보고서 SQL의 날짜·지표·출력 컬럼과 원본 합계를 확인하세요.') from None
    return detail,history,devices


def divide(numerator, denominator):
    return numerator / denominator if denominator else None


def analyze(history, day, comparison_day=None):
    totals = {}
    for d in (comparison_day or day-timedelta(days=7), day):
        selected = [h for h in history if h['report_date'] == d]
        total = {k: int(sum(h[k] for h in selected)) for k in METRICS}
        total['orders_per_visit'] = divide(total['orders'], total['visits'])
        total['average_order_value'] = divide(total['revenue_krw'], total['orders'])
        totals[d.isoformat()] = total
    return totals


def export_excel(path, rows, history, totals, day, source, run_id, analysis=None,device_rows=None):
    wb = Workbook()
    sheet = wb.active
    sheet.title = '일일 보고서'
    prev = day-timedelta(days=7)
    p, c = totals[prev.isoformat()], totals[day.isoformat()]
    sheet.append(['일일 분석 보고서'])
    sheet.append(['데이터', source])
    sheet.append(['보고 대상일', day, '비교일', prev])
    sheet.append(['생성 시각 KST', datetime.now(KST).isoformat(timespec='seconds')])
    sheet.append([])
    sheet.append(['지표', '비교일', '보고일', '변화', '변화 유형'])
    for row_num, (key, label, unit) in enumerate([
        ('visits', '방문 수', '%'), ('orders', '주문 수', '%'),
        ('revenue_krw', '매출 원', '%'), ('orders_per_visit', '주문/방문 비율', 'pp'),
        ('average_order_value', '객단가 원', '%'),
    ], start=7):
        delta = None if p[key] is None or c[key] is None else (
            (c[key]-p[key])*100 if unit == 'pp' else divide(c[key]-p[key], p[key]))
        sheet.append([label, p[key] if p[key] is not None else 'n.a.',
                      c[key] if c[key] is not None else 'n.a.', delta if delta is not None else 'n.a.',
                      '비율 차이' if unit == 'pp' else '상대 증감'])
        for col in (2, 3):
            sheet.cell(row_num, col).number_format = '0.00%' if unit == 'pp' else '#,##0'
        sheet.cell(row_num, 4).number_format = '+0.00"pp";-0.00"pp";0.00"pp"' if unit == 'pp' else '+0.0%;-0.0%;0.0%'
    sheet.append([])
    sheet.append(['채널', '비교일 매출 원', '보고일 매출 원', '증감액 원'])
    for ch in CHANNELS:
        before = next(h['revenue_krw'] for h in history if h['report_date'] == prev and h['channel'] == ch)
        after = next(h['revenue_krw'] for h in history if h['report_date'] == day and h['channel'] == ch)
        sheet.append([ch, before, after, after-before])
        for cell in sheet[sheet.max_row][1:4]:
            cell.number_format = '#,##0;[Red](#,##0);0'
    sheet.append([])
    sheet.append(['계산 기준', '주문/방문 = 총 주문 ÷ 총 방문. 비율 차이는 pp. 0분모는 n.a.'])
    sheet.append(['데이터 범위', 'Adobe 분해 행 합산. API 총계와 대조한 입력이며 기간 순방문자가 아닙니다.'
                 if source.startswith('Adobe Analytics') else '테스트 데이터의 방문은 상호 배타적. 실제 Adobe 총계와는 별도 대조 필요.'])
    sheet.append(['분석 방식', 'SQL·코드 계산. LLM 상태: ' + (analysis or {}).get('status', 'disabled')])
    sheet.append(['실행 ID', run_id])
    sheet['B3'].number_format = sheet['D3'].number_format = 'yyyy-mm-dd'
    detail = wb.create_sheet('조회 데이터')
    detail.append(['날짜', '채널', '디바이스', '방문 수', '페이지 조회 수', '주문 수', '매출 원'])
    for r in rows:
        detail.append([date.fromisoformat(r['date']), *[r[k] for k in FIELDS[1:]]])
    detail.freeze_panes = 'D2'
    detail.auto_filter.ref = detail.dimensions
    for row in detail.iter_rows(min_row=2):
        row[0].number_format = 'yyyy-mm-dd'
        for cell in row[3:]:
            cell.number_format = '#,##0'
    if device_rows is not None:
        device=wb.create_sheet('디바이스 집계');device.append(['날짜','디바이스',*METRICS])
        for r in device_rows:device.append([r['report_date'],r['device'],*[r[k] for k in METRICS]])
        device.freeze_panes='C2';device.auto_filter.ref=device.dimensions
        for row in device.iter_rows(min_row=2):
            row[0].number_format='yyyy-mm-dd'
            for cell in row[2:]:cell.number_format='#,##0'
    for ws, headers in [(sheet, (6, 13)), (detail, (1,))]:
        ws.sheet_view.showGridLines = False
        for row in ws:
            ws.row_dimensions[row[0].row].height = 24
            for cell in row:
                cell.font = Font(name='Arial', size=11)
                cell.alignment = Alignment(vertical='center')
        for n in headers:
            for cell in ws[n]:
                cell.fill = PatternFill('solid', fgColor='213B54')
                cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        for col in 'ABCDEFG':
            ws.column_dimensions[col].width = 22 if col != 'A' else 24
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_options.horizontalCentered = True
    for row in (2, 4, 20, 21, 22, 23):
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        sheet.cell(row, 2).alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[row].height = 32
    sheet['A1'].font = Font(name='Arial', size=17, bold=True)
    sheet.print_area = 'A1:E23'
    sheet.page_setup.fitToHeight = 1
    if analysis:
        ai = wb.create_sheet('AI 분석')
        ai.append(['AI 해석 · 수치 원본은 일일 보고서 참조'])
        ai.append(['상태', analysis['status']])
        ai.append(['공급자 / 모델', f"{analysis['provider']} / {analysis['model']}"])
        ai.append(['안내', 'AI 해석은 검토가 필요합니다. 원인·가설을 사실로 간주하지 마세요.'])
        for line in analysis['text'].splitlines():
            # Force untrusted model text to a string, never an Excel formula.
            ai.append([line])
            ai.cell(ai.max_row, 1).data_type = 's'
        ai.column_dimensions['A'].width = 100
        ai.column_dimensions['B'].width = 65
        for row in ai:
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
            ai.row_dimensions[row[0].row].height = max(30, 16 * (len(str(row[0].value or '')) // 55 + 1))
    wb.save(path)
    wb.close()


def atomic_json(path, obj):
    handle, temporary = tempfile.mkstemp(dir=path.parent, prefix='.json-', suffix='.tmp')
    try:
        with os.fdopen(handle, 'w', encoding='utf-8') as out:
            json.dump(obj, out, ensure_ascii=False, indent=2, default=str)
            out.flush()
            os.fsync(out.fileno())
        os.replace(temporary, path)
    finally:
        Path(temporary).unlink(missing_ok=True)


def run(conn, day, csv_path=None, regenerate=False, report_dir=None, log_dir=None):
    reports = Path(report_dir or os.environ.get('REPORT_DIR', ROOT/'reports'))
    logs = Path(log_dir or os.environ.get('LOG_DIR', ROOT/'logs'))
    reports.mkdir(parents=True, exist_ok=True)
    logs.mkdir(parents=True, exist_ok=True)
    run_id = uuid.uuid4().hex
    record = {'run_id': run_id, 'report_date': day.isoformat(), 'status': 'running',
              'started_at': datetime.now(KST).isoformat(), 'stage': 'lock'}
    record_path = logs/f'{run_id}.json'
    atomic_json(record_path, record)
    locked, temporary = False, None
    try:
        acquire_lock(conn)
        locked = True
        record['stage'] = 'extract'
        days = (day-timedelta(days=7), day)
        rows = extract(days, csv_path)
        fingerprint = hashlib.sha256(json.dumps(rows, sort_keys=True).encode()).hexdigest()
        source = f'CSV 테스트 데이터: {Path(csv_path).name}' if csv_path else (
            {'example_api':'가상 쇼핑몰 HTTP API (example_api)', 'adobe':'Adobe Analytics (채널·디바이스 분해)'}
            .get(os.environ.get('DATA_PROVIDER'), '가상 데이터 (mock)'))
        if not csv_path and os.environ.get('DATA_PROVIDER') == 'adobe':
            source += ' · Report Suite: ' + os.environ.get('ADOBE_REPORT_SUITE_ID', '')
        llm = providers.llm_config()
        import query_workspace as w
        contents=w.pipeline_contents()
        aggregation_sql=sql_statements('aggregate_daily.sql',contents)[0]
        identity = {'llm': llm, 'version': VERSION, 'input_sha256': fingerprint, 'source': source,
                    'database_sha256': hashlib.sha256(f'{conn.host}:{conn.port}:{conn.db}'.encode()).hexdigest(),
                    'sql_sha256': hashlib.sha256(json.dumps(contents,sort_keys=True).encode()).hexdigest(),
                    'aggregation_sql':aggregation_sql}
        manifest_path = reports/f'daily_report_{day}.json'
        if manifest_path.exists() and not regenerate:
            previous = json.loads(manifest_path.read_text(encoding='utf-8'))
            existing = reports/Path(previous['file']).name
            if previous['identity'] != identity:
                raise ValueError('Input, database, SQL or LLM configuration changed. Use --regenerate to create a new report version.')
            if existing.exists() and hashlib.sha256(existing.read_bytes()).hexdigest() == previous['file_sha256']:
                record.update(status='skipped', file=existing.name, stage='complete')
                return existing
        version, output = 1, reports/f'daily_report_{day}.xlsx'
        while output.exists():
            version += 1
            output = reports/f'daily_report_{day}_v{version}.xlsx'
        record['stage'] = 'load_transform'
        atomic_json(record_path, record)
        conn.begin()
        load_transform(conn, rows, days,aggregation_sql,contents)
        rows,history,device_rows=report_data(conn,days,rows,contents)
        totals = analyze(history, day)
        record['stage'] = 'llm'
        atomic_json(record_path, record)
        analysis = providers.summarize(totals, history, llm)
        record['llm_status'] = analysis['status']
        LOG.info('llm: %s', analysis['status'])
        record['stage'] = 'export'
        atomic_json(record_path, record)
        handle, temporary = tempfile.mkstemp(dir=reports, prefix='.report-', suffix='.xlsx')
        os.close(handle)
        export_excel(temporary, rows, history, totals, day, source, run_id, analysis,device_rows)
        # Build before commit: export failure rolls back both raw and history writes.
        conn.commit()
        record['stage'] = 'publish'
        atomic_json(record_path, record)
        # Hard-link publication is atomic and refuses to overwrite an existing file.
        os.link(temporary, output)
        atomic_json(manifest_path, {'run_id': run_id, 'file': output.name, 'identity': identity,
                    'file_sha256': hashlib.sha256(output.read_bytes()).hexdigest(), 'totals': totals,
                    'report_date': day.isoformat(), 'comparison_date': days[0].isoformat(),
                    'analysis': analysis, 'generated_at': datetime.now(KST).isoformat()})
        record.update(status='succeeded', file=output.name, stage='complete')
        LOG.info('complete: %s', output.name)
        return output
    except Exception as error:
        conn.rollback()
        record.update(status='failed', error_type=type(error).__name__)
        # Do not serialize DB exceptions or connection credentials into public reports.
        raise
    finally:
        if temporary:
            Path(temporary).unlink(missing_ok=True)
        record['finished_at'] = datetime.now(KST).isoformat()
        try:
            atomic_json(record_path, record)
        finally:
            if locked:
                with conn.cursor() as cur:
                    cur.execute('SELECT RELEASE_LOCK(%s)', (lock_name(conn),))



def validate_query_plan(plan, day):
    if not isinstance(plan, dict) or set(plan) != {'start', 'end', 'group', 'channel', 'device'}:
        raise ValueError('조회 조건 형식이 올바르지 않습니다. 기간과 채널을 명시해 다시 질문하세요.')
    try:
        start, end = date.fromisoformat(plan['start']), date.fromisoformat(plan['end'])
    except (ValueError, TypeError):
        raise ValueError('조회 날짜를 해석하지 못했습니다.') from None
    if not 0 <= (end-start).days < 31 or end > day or start < date(2000, 1, 1):
        raise ValueError('조회 기간은 기준일 이전의 1~31일이어야 합니다.')
    if (plan['group'] not in ('summary', 'channel', 'device', 'daily') or
            plan['channel'] not in ('', *CHANNELS) or plan['device'] not in ('', *DEVICES)):
        raise ValueError('허용되지 않은 집계 방식 또는 필터입니다.')
    width = (end-start).days+1
    # A single day keeps the existing same-weekday comparison. Ranges use the preceding equal-length period.
    shift = 7 if width == 1 else width
    return start, end, start-timedelta(days=shift), end-timedelta(days=shift)


def ask(conn, day, question):
    """LLM selects bounded parameters, repository templates own all SQL."""
    if not question.strip() or len(question) > 2000:
        raise ValueError('Question must contain 1–2000 characters')
    config = providers.llm_config()
    if config['provider'] == 'none':
        raise ValueError('Set LLM_PROVIDER, LLM_MODEL and the API key in .env first')
    plan = providers.plan_query(question, day)
    start, end, before_start, before_end = validate_query_plan(plan, day)
    dates = [before_start+timedelta(days=i) for i in range((before_end-before_start).days+1)]
    dates += [start+timedelta(days=i) for i in range((end-start).days+1)]
    bounds = (before_start, before_end, start, end)
    coverage_sql = sql_statements('question_coverage.sql')[0]
    sql = sql_statements('question.sql')[('summary', 'channel', 'device', 'daily').index(plan['group'])]
    params = (*bounds, plan['channel'], plan['channel'], plan['device'], plan['device'])
    with conn.cursor() as cur:
        cur.execute('START TRANSACTION READ ONLY')
        try:
            cur.execute(coverage_sql, bounds)
            counts = {r['report_date']: r['row_count'] for r in cur.fetchall()}
            missing = [d.isoformat() for d in dates if counts.get(d) != len(CHANNELS)*len(DEVICES)]
            if missing:
                raise ValueError('Incomplete source: 데이터가 부족한 날짜: ' + ', '.join(missing[:8]) +
                                 (' 외 ' + str(len(missing)-8) + '일' if len(missing)>8 else '') +
                                 '. 해당 날짜의 보고서를 먼저 생성하세요.')
            cur.execute(sql, params)
            rows = cur.fetchall()
        finally:
            conn.rollback()
    grouped = {}
    for row in rows:
        current = start <= row['report_date'] <= end
        offset = (row['report_date']-(start if current else before_start)).days
        label = str(offset+1)+'일차' if plan['group']=='daily' else row['segment']
        segment = grouped.setdefault(label, {'before': {m: 0 for m in METRICS}, 'after': {m: 0 for m in METRICS}})
        for metric in METRICS:
            segment['after' if current else 'before'][metric] += int(row[metric])
    comparisons = {label: {metric: {'before': v['before'][metric], 'after': v['after'][metric],
        'change': v['after'][metric]-v['before'][metric],
        'change_percent': divide((v['after'][metric]-v['before'][metric])*100, v['before'][metric])}
        for metric in METRICS} for label, v in grouped.items()}
    scope = dict(plan, comparison_start=before_start.isoformat(), comparison_end=before_end.isoformat())
    config['prompt'] += ('\n일일 보고서 대신 질문에 직접 답하세요. 제공된 조회 범위와 comparisons만 사용하세요. '
        '합계와 증감률은 코드에서 계산한 값이므로 그대로 사용하고 다시 합산하지 마세요. '
        '하루는 전주 같은 요일, 여러 날은 직전 같은 길이 기간과 비교합니다. '
        '일별 집계의 N일차는 두 기간의 시작일로부터 N번째 날입니다. '
        '없는 지표나 범위 밖 질문은 데이터 부족을 명시하세요. 방문 단위는 건입니다. '
        '[관찰된 변화], [확인할 가설], [다음 분석 제안]을 각각 별도 줄의 제목으로 쓰세요.\n질문: ' + question)
    result = providers.summarize({'scope': scope, 'comparisons': comparisons}, [], config)
    if result['status'] != 'succeeded':
        raise RuntimeError('LLM answer unavailable: ' + result['status'] + ' ' + result.get('error', ''))
    return {'question': question, 'scope': scope, 'dates': [d.isoformat() for d in dates],
            'rows': sum(r['source_rows'] for r in rows), 'sql': sql, 'parameters': [str(p) for p in params],
            'coverage_sql': coverage_sql, 'comparisons': comparisons, 'answer': result['text']}


def main():
    parser = argparse.ArgumentParser(description='MariaDB → SQL → daily Excel local MVP (mock/CSV/example API + optional LLM)')
    commands = parser.add_subparsers(dest='command', required=True)
    commands.add_parser('adobe-check', help='Check Adobe authentication and available metadata; no DB/LLM calls')
    adobe_preview = commands.add_parser('adobe-preview', help='Fetch unmodified Adobe daily breakdowns; no DB/LLM calls')
    adobe_preview.add_argument('--date', type=date.fromisoformat, required=True)
    commands.add_parser('init-db', help='Create two MVP tables in a designated TEST database only')
    report = commands.add_parser('run', help='Generate a new dated Excel report')
    report.add_argument('--date', type=date.fromisoformat,
                        default=datetime.now(KST).date()-timedelta(days=1), help='Report date, default yesterday KST')
    report.add_argument('--csv', type=Path, help='Normalized CSV; overrides DATA_PROVIDER')
    report.add_argument('--regenerate', action='store_true', help='Preserve prior report and publish a new version')
    qa = commands.add_parser('ask', help='Ask using bounded date/filter parameters and read-only SQL templates')
    qa.add_argument('--date', type=date.fromisoformat, required=True)
    qa.add_argument('question')
    commands.add_parser('tables', help='List connected database tables and columns')
    query=commands.add_parser('query',help='Preview a read-only SELECT')
    query.add_argument('--sql',required=True)
    query.add_argument('--date',type=date.fromisoformat,required=True)
    preview=commands.add_parser('table-preview',help='Preview 20 rows of a table in the connected database')
    preview.add_argument('--table',required=True)
    preview.add_argument('--date',type=date.fromisoformat,required=True)
    aggregate=commands.add_parser('aggregation-preview',help='Validate daily aggregation SELECT without DB writes')
    aggregate.add_argument('--sql',required=True)
    aggregate.add_argument('--date',type=date.fromisoformat,required=True)
    stage=commands.add_parser('pipeline-preview',help='Execute a candidate pipeline once and roll back all DB changes')
    stage.add_argument('--name',required=True)
    stage.add_argument('--sql',required=True)
    stage.add_argument('--date',type=date.fromisoformat,required=True)
    for command in ('query-answer','query-report'):
        item=commands.add_parser(command)
        item.add_argument('--id',required=True)
        item.add_argument('--date',type=date.fromisoformat,required=True)
        if command=='query-answer':item.add_argument('question')
        else:
            item.add_argument('--regenerate',action='store_true')
            item.add_argument('--refresh',action='store_true',help='Collect and load daily source data before the saved query')
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format='%(asctime)s %(levelname)s %(message)s')
    try:
        if args.command in ('adobe-check','adobe-preview'):
            import adobe
            client = adobe.Client()
            result = client.check() if args.command == 'adobe-check' else client.collect((args.date-timedelta(days=7),args.date))
            print(json.dumps(result,ensure_ascii=False,indent=2))
            return 0
        with connect() as conn:
            if args.command in ('tables','table-preview','query','query-answer','query-report','aggregation-preview','pipeline-preview'):
                import query_workspace as w
                if args.command=='tables':result=w.schema(conn)
                elif args.command=='table-preview':result=w.preview(conn,args.table,args.date)
                elif args.command=='aggregation-preview':result=w.preview_aggregation(conn,args.sql,args.date)
                elif args.command=='pipeline-preview':result=w.preview_pipeline(conn,args.name,args.sql,args.date)
                elif args.command=='query':result=w.execute(conn,args.sql,w.parameters(args.date))
                elif args.command=='query-answer':result=w.ask(conn,args.date,args.question,args.id)
                else:result={'file':str(w.report(conn,args.date,args.id,args.regenerate,args.refresh))}
                print(json.dumps(result,ensure_ascii=False,default=str))
            elif args.command == 'init-db':
                init_db(conn)
                LOG.info('MVP test schema ready')
            elif args.command == 'ask':
                print(json.dumps(ask(conn, args.date, args.question), ensure_ascii=False, indent=2))
            else:
                print(run(conn, args.date, args.csv, args.regenerate))
    except (ValueError, RuntimeError) as error:
        LOG.error('%s', error)
        return 1
    except Exception as error:
        LOG.error('%s: operation failed; check local DB/service configuration and logs.', type(error).__name__)
        return 1
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
