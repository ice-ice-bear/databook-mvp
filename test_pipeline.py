"""One integration check against the isolated MariaDB demo. Run via Compose."""
import hashlib
import json
import os
from urllib.error import HTTPError, URLError
import tempfile
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import pipeline as p


def expect_failure(call, kind):
    try:
        call()
    except kind:
        return
    raise AssertionError(f'Expected {kind.__name__}')


def check():
    day = date(2026, 9, 27)
    days = (day-timedelta(days=7), day)
    source = p.ROOT/'data/adobe_mock_daily.csv'
    rows = p.extract(days, source)
    assert len(rows) == 30
    # Exercise the actual local HTTP API and its normalization contract.
    with patch.dict(os.environ, {'DATA_PROVIDER': 'example_api'}):
        assert p.extract(days) == p.validate(p.mock_rows(days), days)
    assert p.divide(1, 0) is None
    expect_failure(lambda: p.validate(rows[:-1], days), ValueError)
    expect_failure(lambda: p.validate(rows+[rows[0]], days), ValueError)
    expect_failure(lambda: p.validate([dict(rows[0], visits=-1), *rows[1:]], days), ValueError)
    expect_failure(lambda: p.extract((date(2025, 1, 1), date(2025, 1, 8)), source), ValueError)
    with tempfile.TemporaryDirectory() as tmp, p.connect() as conn:
        root = Path(tmp)
        p.init_db(conn)
        kwargs = dict(csv_path=source, report_dir=root/'reports', log_dir=root/'logs')
        output = p.run(conn, day, **kwargs)
        digest = hashlib.sha256(output.read_bytes()).hexdigest()
        wb = load_workbook(output, data_only=True)
        assert wb['일일 보고서']['B9'].value == 24079860
        assert wb['일일 보고서']['C9'].value == 22273572
        assert wb['조회 데이터'].max_row == 31
        assert wb['디바이스 집계'].max_row == 7
        assert wb['일일 보고서']['B3'].value.date() == day
        wb.close()
        with conn.cursor() as cur:
            cur.execute('SELECT COUNT(*) AS n FROM mvp_raw_daily WHERE report_date IN (%s,%s)', days)
            assert cur.fetchone()['n'] == 30
            cur.execute('SELECT COUNT(*) AS n FROM mvp_history_daily WHERE report_date IN (%s,%s)', days)
            assert cur.fetchone()['n'] == 10
        conn.commit()
        assert p.run(conn, day, **kwargs) == output
        assert len(list((root/'reports').glob('*.xlsx'))) == 1
        regenerated = p.run(conn, day, regenerate=True, **kwargs)
        assert regenerated.stem.endswith('_v2')
        assert hashlib.sha256(output.read_bytes()).hexdigest() == digest
        assert p.run(conn, day, **kwargs) == regenerated
        with patch.object(conn, 'host', 'another-database-host'):
            expect_failure(lambda: p.run(conn, day, **kwargs), ValueError)
        with p.connect() as second:
            p.acquire_lock(second)
            expect_failure(lambda: p.run(conn, day, **kwargs), RuntimeError)
        # Changed input must not silently reuse an earlier report.
        expect_failure(lambda: p.run(conn, day, report_dir=root/'reports', log_dir=root/'logs'), ValueError)
        with conn.cursor() as cur:
            cur.execute('SELECT * FROM mvp_raw_daily WHERE report_date IN (%s,%s) ORDER BY report_date,channel,device', days)
            before = cur.fetchall()
        conn.commit()
        with patch.object(p, 'export_excel', side_effect=OSError('simulated disk full')):
            expect_failure(lambda: p.run(conn, day, regenerate=True, report_dir=root/'reports', log_dir=root/'logs'), OSError)
        with conn.cursor() as cur:
            cur.execute('SELECT * FROM mvp_raw_daily WHERE report_date IN (%s,%s) ORDER BY report_date,channel,device', days)
            assert cur.fetchall() == before, 'Failed export must roll back changed raw input'
        conn.commit()
        assert not (root/'reports'/f'daily_report_{day}_v3.xlsx').exists()
        assert hashlib.sha256(output.read_bytes()).hexdigest() == digest
        states = [json.loads(f.read_text(encoding='utf-8'))['status'] for f in (root/'logs').glob('*.json')]
        assert {'succeeded', 'failed', 'skipped'} <= set(states)
        # Zero denominator should be a visible n.a., not a spreadsheet error or false zero.
        zero = [dict(r, visits=0, page_views=0, orders=0, revenue_krw=0) for r in rows]
        p.validate(zero, days)
        conn.begin()
        history = p.load_transform(conn, zero, days)
        totals = p.analyze(history, day)
        assert totals[day.isoformat()]['orders_per_visit'] is None
        p.export_excel(root/'zero.xlsx', zero, history, totals, day, 'zero test', 'test')
        conn.rollback()
        wb = load_workbook(root/'zero.xlsx', data_only=True)
        assert wb['일일 보고서']['C10'].value == 'n.a.'
        assert wb['일일 보고서']['D9'].value == 'n.a.'
        wb.close()
        # Both LLM adapters: payload/auth contract, text extraction and safe failure.
        import providers as api
        for provider in ('openai', 'gemini'):
            response = ({'status': 'completed', 'output': [{'type': 'message', 'content':
                        [{'type': 'output_text', 'text': '=1+1'}]}]} if provider == 'openai' else
                        {'candidates': [{'finishReason': 'STOP', 'content': {'parts': [{'text': '=1+1'}]}}]})
            with patch.dict(os.environ, {'LLM_PROVIDER': provider, 'LLM_MODEL': 'test-model',
                                         'OPENAI_API_KEY': 'test-only', 'GEMINI_API_KEY': 'test-only'}):
                with patch.object(api, 'request_json', return_value=response) as request:
                    result = api.summarize(totals, history, api.llm_config())
                    assert result['status'] == 'succeeded' and result['text'] == '=1+1'
                    url, payload, headers = request.call_args.args
                    assert 'test-only' not in url
                    assert ('Authorization' if provider == 'openai' else 'x-goog-api-key') in headers
                    assert payload['model'] == 'test-model' if provider == 'openai' else '/test-model:' in url
                with patch.object(api, 'request_json', side_effect=HTTPError('hidden', 401, 'secret', {}, None)):
                    failed = p.run(conn, day, regenerate=True, **kwargs)
                    manifest = json.loads((root/'reports'/f'daily_report_{day}.json').read_text(encoding='utf-8'))
                    assert manifest['analysis']['status'] == 'failed'
                    assert manifest['analysis']['error'] == 'HTTP_401'
                    assert 'test-only' not in json.dumps(manifest)
                    wb = load_workbook(failed)
                    assert wb['AI 분석']['B2'].value == 'failed'
                    wb.close()
                with patch.object(api, 'request_json', return_value=response):
                    analyzed = p.run(conn, day, regenerate=True, **kwargs)
                    wb = load_workbook(analyzed)
                    assert wb['AI 분석']['A5'].value == '=1+1'
                    assert wb['AI 분석']['A5'].data_type == 's'
                    wb.close()
                for invalid in ({}, [], {'status': 'incomplete'}, {'candidates': [{'finishReason': 'MAX_TOKENS'}]}):
                    with patch.object(api, 'request_json', return_value=invalid):
                        assert api.summarize(totals, history, api.llm_config())['status'] == 'failed'
                with patch.object(api, 'request_json', side_effect=URLError('secret')):
                    assert api.summarize(totals, history, api.llm_config())['status'] == 'failed'
                with patch.dict(os.environ, {'OPENAI_API_KEY': '', 'GEMINI_API_KEY': ''}), patch.object(api, 'request_json') as request:
                    assert api.summarize(totals, history, api.llm_config())['status'] == 'not_configured'
                    request.assert_not_called()
        with patch.dict(os.environ, {'LLM_PROVIDER': 'openai', 'LLM_MODEL': 'test-model'}):
            plan = {'start': day.isoformat(), 'end': day.isoformat(), 'group': 'device', 'channel': '', 'device': ''}
            with patch.object(p.providers, 'plan_query', return_value=plan), patch.object(p.providers, 'summarize', return_value={'status': 'succeeded', 'text': 'test answer'}) as model:
                answer = p.ask(conn, day, '모바일 비교')
                assert answer['rows'] == 30 and answer['answer'] == 'test answer'
                assert model.call_args.args[1] == []  # Only precomputed aggregates go to the LLM.
                comparisons = answer['comparisons']
                for device in p.DEVICES:
                    expected = sum(row['revenue_krw'] for row in rows if row['date'] == day.isoformat() and row['device'] == device)
                    assert comparisons[device]['revenue_krw']['after'] == expected
                assert sum(x['revenue_krw']['after'] for x in comparisons.values()) == 22273572
                expect_failure(lambda: p.ask(conn, day, ' '), ValueError)
                for key, value in [('group', 'channel; DROP TABLE mvp_raw_daily'), ('channel', "' OR 1=1"), ('start', '2026-01-01'), ('end', '2026-09-28')]:
                    expect_failure(lambda key=key,value=value: p.validate_query_plan(dict(plan, **{key:value}), day), ValueError)
                plan.update(start='2025-01-08', end='2025-01-08')
                expect_failure(lambda: p.ask(conn, day, '비교'), ValueError)
                # Four complete dates: two-day period against the preceding two days.
                for offset in (1, 0):
                    pair=(day-timedelta(days=offset+2), day-timedelta(days=offset))
                    p.load_transform(conn, p.mock_rows(pair), pair)
                conn.commit()
                plan.update(start=(day-timedelta(days=1)).isoformat(), end=day.isoformat(), group='channel', channel='paid_search', device='mobile')
                answer=p.ask(conn, day, '검색 모바일 2일 매출')
                assert answer['rows']==4 and list(answer['comparisons'])==['paid_search']
                expected=sum(r['revenue_krw'] for r in p.mock_rows((day-timedelta(days=1),day)) if r['channel']=='paid_search' and r['device']=='mobile')
                assert answer['comparisons']['paid_search']['revenue_krw']['after']==expected
                assert answer['scope']['comparison_start']==(day-timedelta(days=3)).isoformat()
                plan['group']='daily'
                daily=p.ask(conn,day,'일별 추세')
                assert set(daily['comparisons'])=={'1일차','2일차'}
                assert sum(r['revenue_krw']['after'] for r in daily['comparisons'].values())==expected
                plan.update(group='summary',channel='',device='')
                assert list(p.ask(conn,day,'전체 매출')['comparisons'])==['전체']
            with patch.object(p.providers, 'summarize', return_value={'status':'succeeded','text':'SQL please'}):
                expect_failure(lambda: p.providers.plan_query('질문',day),ValueError)
            with patch.object(p.providers, 'summarize', return_value={'status':'succeeded','text':json.dumps(plan)}):
                assert p.providers.plan_query('질문',day)==plan
    print('PASS: HTTP provider, OpenAI/Gemini contracts, LLM fallback, formula safety; ')
    print('PASS: MariaDB load/SQL/export; known totals; duplicate/missing/invalid input; '
          'idempotency; versions; concurrent lock; changed input; rollback; zero denominator.')


if __name__ == '__main__':
    check()
